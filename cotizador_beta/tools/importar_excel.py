"""
Importador de la planilla "Cotizador de Aberturas.xlsx".

Lee la hoja *Precios Base* y actualiza (o crea) en la base:

    * Líneas de aluminio y sus colores, con el precio $/m² de perfil
    * Tipos de vidrio con su precio $/m²
    * Tipologías declaradas en la planilla
    * Parámetros de mano de obra, instalación, accesorios, mosquitero,
      premarco, margen e IVA

Es tolerante con el layout: busca los encabezados de cada bloque por texto en
lugar de asumir números de fila fijos, así sigue funcionando si agregás filas.

Requiere ``openpyxl`` (``pip install openpyxl``).
"""

from __future__ import annotations

import unicodedata
from pathlib import Path


def _normalizar(texto) -> str:
    """Minúsculas, sin acentos y sin superíndices.

    Se usa NFKD (no NFD) a propósito: las etiquetas de la planilla dicen "($/m²)"
    y sólo la descomposición de compatibilidad convierte '²' en '2'.
    """
    if texto is None:
        return ""
    texto = str(texto).strip().lower()
    return "".join(c for c in unicodedata.normalize("NFKD", texto)
                   if unicodedata.category(c) != "Mn")


# Etiqueta en la planilla -> (clave del parámetro, es_porcentaje)
MAPA_PARAMETROS = {
    "mano de obra fabricacion ($/m2)": ("mo_por_m2_directo", False),
    "instalacion / colocacion ($/m2)": ("log_colocacion_m2", False),
    "accesorios promedio ($/hoja)": ("accesorios_por_hoja", False),
    "mosquitero ($/m2)": ("mosquitero_m2", False),
    "premarco ($/m2)": ("premarco_m2", False),
    "margen de ganancia (%)": ("margen_pct", True),
    "iva (%)": ("iva_pct", True),
}

ENCABEZADOS = {
    "lineas": "lineas de aluminio",
    "vidrios": "tipos de vidrio",
    "tipologias": "tipos de abertura",
    "otros": "otros costos",
}


def _leer_hoja(ws) -> list[list]:
    return [[c.value for c in fila] for fila in ws.iter_rows()]


def _ubicar_bloques(filas: list[list]) -> dict[str, int]:
    """Devuelve el índice de fila donde empieza cada bloque."""
    posiciones = {}
    for i, fila in enumerate(filas):
        texto = _normalizar(fila[0] if fila else "")
        for clave, patron in ENCABEZADOS.items():
            if texto.startswith(patron):
                posiciones[clave] = i
    return posiciones


def importar(db, ruta_excel: str | Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # el llamador muestra el mensaje al usuario
        raise ImportError("openpyxl no está instalado") from exc

    wb = load_workbook(Path(ruta_excel), data_only=True, read_only=True)
    hoja = None
    for nombre in wb.sheetnames:
        if "precio" in _normalizar(nombre):
            hoja = wb[nombre]
            break
    if hoja is None:
        hoja = wb[wb.sheetnames[0]]

    filas = _leer_hoja(hoja)
    bloques = _ubicar_bloques(filas)

    contadores = {"lineas": 0, "colores": 0, "vidrios": 0, "tipologias": 0, "parametros": 0}

    # ------------------------------------------------------------------ líneas
    inicio = bloques.get("lineas")
    if inicio is not None:
        for fila in filas[inicio + 2:]:
            nombre = fila[0] if fila else None
            if not nombre or not str(nombre).strip():
                break
            color = str(fila[1] or "Natural").strip()
            try:
                precio_m2 = float(fila[2] or 0)
            except (TypeError, ValueError):
                continue

            nombre = str(nombre).strip()
            linea = db.query_one("SELECT * FROM lineas WHERE nombre = ?", (nombre,))
            if linea is None:
                lid = db.insertar("lineas", {"nombre": nombre, "descripcion": "Importada de Excel",
                                             "modo_costeo": "m2", "activo": 1})
                contadores["lineas"] += 1
            else:
                lid = linea["id"]

            existente = db.precio_linea(lid, color)
            if existente:
                db.actualizar("linea_precios", existente["id"], {"precio_m2_perfil": precio_m2})
            else:
                db.insertar("linea_precios", {"linea_id": lid, "color": color, "precio_kg": 0,
                                              "precio_m2_perfil": precio_m2, "activo": 1})
            contadores["colores"] += 1

    # ----------------------------------------------------------------- vidrios
    inicio = bloques.get("vidrios")
    if inicio is not None:
        for fila in filas[inicio + 2:]:
            nombre = fila[0] if fila else None
            if not nombre or not str(nombre).strip():
                break
            try:
                precio = float(fila[1] or 0)
            except (TypeError, ValueError):
                continue
            nombre = str(nombre).strip()
            existente = db.query_one("SELECT id FROM vidrios WHERE nombre = ?", (nombre,))
            if existente:
                db.actualizar("vidrios", existente["id"], {"precio_m2": precio})
            else:
                tipo = ("DVH" if "dvh" in _normalizar(nombre)
                        else "Laminado" if "lamin" in _normalizar(nombre) else "Float")
                db.insertar("vidrios", {"nombre": nombre, "tipo": tipo, "espesor_mm": 4,
                                        "precio_m2": precio, "plancha_ancho_mm": 3600,
                                        "plancha_alto_mm": 2500,
                                        "precio_plancha": round(precio * 9.0, 2),
                                        "desperdicio_pct": 0.10, "activo": 1})
            contadores["vidrios"] += 1

    # -------------------------------------------------------------- tipologías
    nuevas_tipologias: list[str] = []
    inicio = bloques.get("tipologias")
    if inicio is not None:
        existentes = [_normalizar(t["nombre"]) for t in db.tipologias(solo_activas=False)]
        for fila in filas[inicio + 1:]:
            nombre = fila[0] if fila else None
            if not nombre or not str(nombre).strip():
                break
            nombre = str(nombre).strip()
            # "Ventana Corrediza" (planilla) ya está cubierta por
            # "Ventana Corrediza 2 hojas" (base): se compara por prefijo en
            # ambos sentidos para no duplicar tipologías equivalentes.
            clave = _normalizar(nombre)
            if any(e.startswith(clave) or clave.startswith(e) for e in existentes):
                continue
            codigo = "".join(p[0] for p in _normalizar(nombre).split())[:5].upper() or "TIP"
            sufijo = 1
            base = codigo
            while db.query_one("SELECT id FROM tipologias WHERE codigo = ?", (codigo,)):
                codigo = f"{base}{sufijo}"
                sufijo += 1
            n = _normalizar(nombre)
            esquema = ("corrediza" if "corrediza" in n
                       else "fijo" if "fijo" in n
                       else "banderola" if "banderola" in n
                       else "puerta_batiente" if "puerta" in n
                       else "batiente")
            db.insertar("tipologias", {"codigo": codigo, "nombre": nombre,
                                       "hojas_default": 2 if "corrediza" in n else 1,
                                       "admite_mosquitero": 0 if "fijo" in n else 1,
                                       "admite_premarco": 1, "esquema": esquema,
                                       "horas_por_m2": 0.85, "activo": 1})
            existentes.append(clave)
            nuevas_tipologias.append(f"{codigo} — {nombre}")
            contadores["tipologias"] += 1

    # ------------------------------------------------------- otros parámetros
    inicio = bloques.get("otros")
    if inicio is not None:
        for fila in filas[inicio + 1:]:
            etiqueta = _normalizar(fila[0] if fila else "")
            if not etiqueta:
                break
            destino = MAPA_PARAMETROS.get(etiqueta)
            if not destino:
                continue
            clave, es_pct = destino
            try:
                valor = float(fila[1] or 0)
            except (TypeError, ValueError):
                continue
            if es_pct and valor > 1:
                valor = valor / 100.0
            db.set_parametro(clave, valor)
            contadores["parametros"] += 1

    wb.close()

    resumen = (
        f"Importación desde {Path(ruta_excel).name}:\n"
        f"  Líneas nuevas: {contadores['lineas']}\n"
        f"  Colores/precios actualizados: {contadores['colores']}\n"
        f"  Vidrios: {contadores['vidrios']}\n"
        f"  Tipologías nuevas: {contadores['tipologias']}\n"
        f"  Parámetros: {contadores['parametros']}\n"
    )
    if nuevas_tipologias:
        resumen += (
            "\nEstas tipologías se crearon vacías y todavía NO tienen fórmulas de "
            "despiece; cargalas en Materiales > Fórmulas de despiece:\n  - "
            + "\n  - ".join(nuevas_tipologias) + "\n"
        )
    resumen += (
        "\nNota: la planilla trae el precio del aluminio por m² de abertura. "
        "Las líneas importadas quedan en modo de costeo 'm2'. "
        "Para cotizar por despiece cargá el precio $/kg y pasalas a modo 'kg'."
    )
    return resumen
